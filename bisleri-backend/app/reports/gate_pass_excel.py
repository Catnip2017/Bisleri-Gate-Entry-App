# app/reports/gate_pass_excel.py — builds the downloadable .xlsx for the
# Gate Pass "Reports" tab (added 1 Oct 2026).
#
# Design notes:
#  - Summary sheet is always first (per spec) — a plain label/value sheet,
#    not a table (it's tiles, not tabular data).
#  - Register/Detail/Department Summary sheets are built as real Excel
#    Table objects (openpyxl.worksheet.table.Table), not just an
#    auto_filter range — Excel shows the filter-dropdown arrows on every
#    header cell and banded rows the moment the file opens, no manual
#    "Format as Table" step needed on the user's end.
#  - Values are written as plain numbers/text (not formulas): this is a
#    point-in-time export of live DB data, not an editable template — the
#    numbers themselves came from the DB query, so there's nothing for a
#    formula to recompute.
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
TITLE_FONT = Font(name=FONT_NAME, size=14, bold=True, color="1F4E78")
SUBTITLE_FONT = Font(name=FONT_NAME, size=10, italic=True, color="595959")
NOTE_FONT = Font(name=FONT_NAME, size=9, italic=True, color="7F7F7F")
BODY_FONT = Font(name=FONT_NAME, size=10)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

TABLE_STYLE = TableStyleInfo(
    name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
    showRowStripes=True, showColumnStripes=False,
)


def _title_block(ws, span, title, subtitle):
    ws.merge_cells(f"A1:{span}1")
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws.merge_cells(f"A2:{span}2")
    ws["A2"] = subtitle
    ws["A2"].font = SUBTITLE_FONT


def _write_table(ws, headers, rows, start_row, col_widths, table_name,
                  date_cols=(), currency_cols=(), bool_cols=()):
    """Writes a header row + data rows styled for an Excel Table, registers
    the Table object (built-in filter dropdowns), freezes the header, and
    returns the last data row number."""
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[start_row].height = 30

    r = start_row + 1
    for row in rows:
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = BODY_FONT
            cell.border = BORDER
            if c in date_cols and val is not None:
                cell.number_format = "dd-mmm-yyyy"
            if c in currency_cols:
                cell.number_format = "#,##0.00;(#,##0.00);-"
        r += 1
    last_row = max(r - 1, start_row + 1)   # Table needs at least one data row

    if r - 1 < start_row + 1:
        # No data rows at all — add one blank row so the Table range is valid.
        for c in range(1, len(headers) + 1):
            ws.cell(row=start_row + 1, column=c)
        last_row = start_row + 1

    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1).coordinate

    last_col = get_column_letter(len(headers))
    table_ref = f"A{start_row}:{last_col}{last_row}"
    tbl = Table(displayName=table_name, ref=table_ref)
    tbl.tableStyleInfo = TABLE_STYLE
    ws.add_table(tbl)
    return last_row


def _summary_rows(ws, pairs, start_row=4):
    r = start_row
    for label, value, is_currency in pairs:
        if label == "":
            r += 1
            continue
        lc = ws.cell(row=r, column=1, value=label)
        lc.font = BODY_FONT
        vc = ws.cell(row=r, column=2, value=value)
        vc.font = Font(name=FONT_NAME, size=10, bold=True)
        vc.alignment = Alignment(horizontal="center")
        if is_currency:
            vc.number_format = "₹#,##0;(₹#,##0);-"
        r += 1
    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 16


def build_creator_workbook(fy_label: str, summary: dict, rows: list, generated_note: str) -> io.BytesIO:
    wb = Workbook()

    sm = wb.active
    sm.title = "Summary"
    _title_block(sm, "B", f"Gate Pass Creator / IT Admin — FY {fy_label} Summary",
                 generated_note)
    sc = summary["status_counts"]
    pairs = [
        ("Total Gate Passes Raised", summary["total_passes"], False),
        ("Returnable (R) Passes", summary["returnable_count"], False),
        ("Non-Returnable (NR) Passes", summary["non_returnable_count"], False),
        ("", "", False),
        ("Status — Open", sc.get("Open", 0), False),
        ("Status — Released", sc.get("Released", 0), False),
        ("Status — Dispatched", sc.get("Dispatched", 0), False),
        ("Status — Partially Received", sc.get("Partially Received", 0), False),
        ("Status — Inward Received", sc.get("Inward Received", 0), False),
        ("Status — Cancelled", sc.get("Cancelled", 0), False),
        ("Status — Closed Without Return", sc.get("Closed Without Return", 0), False),
        ("", "", False),
        ("R Passes Pending Return (dispatched, not yet received)", summary["pending_return_count"], False),
        ("...of which Overdue (past Expected Inward Date)", summary["overdue_count"], False),
        ("Total Amount — Chargeable lines (₹)", summary["total_chargeable_amount"], True),
    ]
    _summary_rows(sm, pairs)

    reg = wb.create_sheet("Register")
    _title_block(reg, "S", f"Gate Pass Creator / IT Admin — FY {fy_label} Register", generated_note)
    headers = [
        "Gate Pass No.", "Pass Type", "Status", "Location", "Department",
        "Vendor/Customer", "Created By", "Document Date", "Dispatch Date/Time",
        "Expected Inward Date", "Actual Inward Date", "Days Outstanding",
        "Line Count", "Total Qty", "Total Amount", "Chargeable?",
        "Cancelled By / Reason", "Remarks",
    ]
    data = []
    for row in rows:
        data.append([
            row["gate_pass_no"], row["pass_type"], row["status"], row["location_code"],
            row["department"], row["party_name"], row["created_by"], row["document_date"],
            row["dispatched_at"], row["expected_inward_date"], row["actual_inward_date"],
            row["days_outstanding"] if row["days_outstanding"] is not None else "-",
            row["line_count"], row["total_qty"], row["total_amount"],
            "Yes" if row["chargeable"] else "No", row["cancel_info"] or "", row["remarks"] or "",
        ])
    col_widths = [18, 9, 13, 9, 12, 22, 15, 13, 16, 15, 16, 13, 9, 9, 12, 10, 26, 28]
    _write_table(reg, headers, data, start_row=4, col_widths=col_widths,
                 table_name="CreatorRegister", date_cols=(8, 9, 10, 11), currency_cols=(15,))

    wb.active = 0  # Summary stays the sheet Excel opens on
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_dispatcher_workbook(fy_label: str, summary: dict, rows: list,
                               dept_summary: list, generated_note: str) -> io.BytesIO:
    wb = Workbook()

    sm = wb.active
    sm.title = "Summary"
    _title_block(sm, "B", f"Gate Pass Dispatcher / Security — FY {fy_label} Item Reconciliation Summary",
                 generated_note)
    pairs = [
        ("Total Item/Asset Lines Dispatched", summary["total_lines"], False),
        ("Total Qty Sent", summary["total_qty_sent"], False),
        ("Departments Covered", summary["departments"], False),
        ("Lines Pending Return", summary["pending_lines"], False),
        ("...of which Overdue (past Expected Return Date)", summary["overdue_lines"], False),
    ]
    _summary_rows(sm, pairs)

    det = wb.create_sheet("Detail")
    _title_block(det, "P", f"Gate Pass Dispatcher / Security — FY {fy_label} Item Reconciliation — Detail",
                 generated_note)
    headers = [
        "Department", "Gate Pass No.", "Pass Type", "Line Type", "Item Code / Asset No.",
        "Description of Goods", "Qty Sent", "Unit", "Dispatch Date", "Expected Return Date",
        "Qty Returned", "Actual Return Date", "Reconciliation Status", "Days Outstanding",
        "Vendor/Customer", "Remarks",
    ]
    data = []
    for row in rows:
        data.append([
            row["department"], row["gate_pass_no"], row["pass_type"], row["line_type"],
            row["item_code"] or "", row["description"], row["qty_sent"], row["unit"],
            row["dispatch_date"], row["expected_return_date"], row["qty_returned"],
            row["actual_return_date"], row["reconciliation_status"],
            row["days_outstanding"] if row["days_outstanding"] is not None else "-",
            row["party_name"], row["remarks"] or "",
        ])
    col_widths = [12, 18, 9, 11, 20, 26, 9, 7, 13, 15, 11, 15, 17, 13, 22, 26]
    _write_table(det, headers, data, start_row=4, col_widths=col_widths,
                 table_name="DispatcherDetail", date_cols=(9, 10, 12))

    dp = wb.create_sheet("Department Summary")
    _title_block(dp, "I", f"FY {fy_label} — Department-wise Item Reconciliation Summary", generated_note)
    dp_headers = [
        "Department", "Total Lines Sent", "Total Qty Sent", "Lines Fully Returned",
        "Lines Partially Returned", "Lines Pending", "NR Lines (outbound only)",
        "Oldest Pending (days)", "Overdue Lines (R, pending, past due date)",
    ]
    dp_data = []
    for d in dept_summary:
        dp_data.append([
            d["department"], d["total_lines"], d["total_qty_sent"], d["lines_fully_returned"],
            d["lines_partially_returned"], d["lines_pending"], d["nr_lines"],
            d["oldest_pending_days"] if d["oldest_pending_days"] is not None else "-",
            d["overdue_lines"],
        ])
    dp_col_widths = [14, 14, 13, 16, 18, 12, 18, 16, 22]
    _write_table(dp, dp_headers, dp_data, start_row=4, col_widths=dp_col_widths,
                 table_name="DepartmentSummary")

    wb.active = 0  # Summary stays the sheet Excel opens on
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
